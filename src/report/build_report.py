"""Rebuild the final report from the code: outputs/report/FBAR2_2026_KTQT_FTUers_rebuilt.docx (and .pdf).

The report is assembled from a Word template (src/report/template/base_report.docx) that holds the styles,
the appendix text and the reference list. This script
  1. replaces Sections 1 to 5 and the executive summary with src/report/content.py,
  2. rebuilds the table of contents and the list of figures and tables,
  3. applies the appendix changes (Tables E3, F5, F6, G1, Figure F2, renumbering, new references),
  4. swaps every figure for the version drawn by src/figures (outputs/figures),
  5. drops references that are no longer cited and builds the title page.

Usage:
    python src/report/build_report.py              # docx only, page numbers from src/report/pages.json
    python src/report/build_report.py --pdf        # also export a PDF (needs LibreOffice)
    python src/report/build_report.py --paginate   # recompute page numbers with LibreOffice, then build
"""
import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Emu
from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
import analysis as AN  # noqa: E402
from config import FIGURES, OUT_REPORT, ROOT  # noqa: E402
from content import BODY, EXEC  # noqa: E402
from figures.style import PRINT_SIZE  # noqa: E402
from title_page import build_title_page  # noqa: E402

TEMPLATE = HERE / "template" / "base_report.docx"
OUT_DOCX = OUT_REPORT / "FBAR2_2026_KTQT_FTUers_rebuilt.docx"
PAGES = HERE / "pages.json"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
WP_NS = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

APPENDICES = ["Appendices", "Appendix A. Data preparation and quality", "Appendix B. DNSE financial data",
              "Appendix C. Market and peer data", "Appendix D. Customer review analysis",
              "Appendix E. Findex segment analysis", "Appendix F. Assumptions, sensitivity and alternative explanations",
              "Appendix G. Regulatory basis and technology proposals", "References"]
APPENDIX_FIGURES = {"Figure B1.": "figB1_revenue_profit.png", "Figure B2.": "figB2_revenue_allocation.png",
                    "Figure B3.": "figB3_yields_funding.png", "Figure B4.": "figB4_asset_composition.png",
                    "Figure C1.": "figC1_peer_per_account.png", "Figure D1.": "figD1_rating_by_year.png",
                    "Figure D2.": "figD2_store_ratings.png", "Figure E1.": "figE1_findex_behaviour.png",
                    "Figure F1.": "figF1_scenarios.png"}


# ------------------------------------------------------------------ XML helpers
def E(tag, **attrs):
    e = etree.Element(W + tag)
    for k, val in attrs.items():
        e.set(W + k, str(val))
    return e


def run(text, bold=False, sz=None):
    r = E("r")
    rp = E("rPr")
    if bold:
        rp.append(E("b")); rp.append(E("bCs"))
    if sz:
        rp.append(E("sz", val=sz)); rp.append(E("szCs", val=sz))
    if len(rp):
        r.append(rp)
    t = E("t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def text_of(e):
    return "".join(t.text or "" for t in e.iter(W + "t"))


class Builder:
    def __init__(self, pages):
        self.doc = Document(str(TEMPLATE))
        self.body = self.doc.element.body
        self.tpl = list(self.body.iterchildren())      # template elements, in template order
        self.pages = pages

    # -------------------------------------------------------------- paragraph factories (template formats)
    def para_from(self, tpl, parts):
        p = E("p")
        p.append(copy.deepcopy(tpl.find(W + "pPr")))
        for txt, b, sz in parts:
            p.append(run(txt, b, sz))
        return p

    def rich(self, text):
        parts = [(seg, i % 2 == 1, None) for i, seg in enumerate(re.split(r"\*\*", text)) if seg]
        return self.para_from(self.tpl[83], parts)

    def heading(self, text, level):
        h1 = {"1.": 64, "2.": 79, "3.": 116, "4.": 140, "5.": 159}
        p = copy.deepcopy(self.tpl[h1[text[:2]] if level == 1 else 65])
        rs = p.findall(W + "r")
        for r in rs[1:]:
            p.remove(r)
        rs[0].find(W + "t").text = text
        if level == 1 and text[:2] in ("3.", "4."):
            pb = p.find(W + "pPr").find(W + "pageBreakBefore")
            if pb is not None:
                p.find(W + "pPr").remove(pb)
        return p

    def fig_caption(self, label, title, source):
        p = E("p")
        p.append(copy.deepcopy(self.tpl[85].find(W + "pPr")))
        p.append(run(f"{label} ", True, 19))
        p.append(run(title, False, 19))
        r = E("r"); r.append(E("br")); p.append(r)
        p.append(run(source, False, 17))
        return p

    def table_caption(self, label, title):
        return self.para_from(self.tpl[125], [(f"{label} ", True, 19), (title, False, 19)])

    def source_line(self, text):
        return self.para_from(self.tpl[127], [(text, False, 17)])

    def box(self, label, text):
        return self.para_from(self.tpl[89], [(label + " ", True, None), (text, False, None)])

    def picture(self, png):
        cx, cy = PRINT_SIZE[png.name]                  # printed size, as in the submitted report
        p = self.doc.add_paragraph()
        p.add_run().add_picture(str(png), width=Emu(cx), height=Emu(cy))
        el = p._p
        el.getparent().remove(el)
        el.insert(0, copy.deepcopy(self.tpl[143].find(W + "pPr")))
        return el

    def table(self, header, rows, widths, align_numbers=True):
        tt = self.tpl[126]
        t = E("tbl")
        tp = copy.deepcopy(tt.find(W + "tblPr"))
        tp.find(W + "tblW").set(W + "w", str(sum(widths)))
        t.append(tp)
        g = E("tblGrid")
        for w in widths:
            g.append(E("gridCol", w=w))
        t.append(g)
        trs = tt.findall(W + "tr")

        def row(cells, hdr):
            src = trs[0] if hdr else trs[1]
            tr = E("tr")
            tr.append(copy.deepcopy(src.find(W + "trPr")))
            tc_t = src.findall(W + "tc")[0]
            for i, (txt, w) in enumerate(zip(cells, widths)):
                tc = copy.deepcopy(tc_t)
                tc.find(W + "tcPr").find(W + "tcW").set(W + "w", str(w))
                p = tc.find(W + "p")
                for r in p.findall(W + "r"):
                    p.remove(r)
                jc = p.find(W + "pPr").find(W + "jc")
                if jc is not None and align_numbers:
                    numeric = re.match(r"^[\d.,%() =stndrh-]+$", str(txt))
                    jc.set(W + "val", "center" if (i and numeric) else "left")
                p.append(run(str(txt), hdr, 18))
                tr.append(tc)
            return tr

        t.append(row(header, True))
        for r in rows:
            t.append(row(r, False))
        return t

    def find(self, start, tag="p"):
        for e in self.body.iterchildren():
            if e.tag == W + tag and text_of(e).startswith(start):
                return e
        raise KeyError(start)

    def replace_in(self, e, old, new):
        hit = False
        for t in e.iter(W + "t"):
            if t.text and old in t.text:
                t.text = t.text.replace(old, new)
                hit = True
        if not hit:
            raise KeyError(f"{old!r} not found")

    def appendix_elements(self):
        els = list(self.body.iterchildren())
        start = next(i for i, e in enumerate(els) if e.tag == W + "p" and text_of(e).strip() == "Appendices"
                     and e.xpath("./w:pPr/w:pStyle/@w:val") == ["Heading1"])
        return els[start:]

    def references_heading(self):
        for e in self.body.iterchildren():
            if e.tag == W + "p" and text_of(e).strip().startswith("References") and \
                    e.xpath("./w:pPr/w:pStyle/@w:val") == ["Heading1"]:
                return e
        raise KeyError("References heading")

    # -------------------------------------------------------------- 1. main body and front matter
    def main_body(self):
        new = []
        for b in BODY:
            k = b[0]
            if k in ("h1", "h2"):
                new.append(self.heading(b[1], 1 if k == "h1" else 2))
            elif k == "p":
                new.append(self.rich(b[1]))
            elif k == "box":
                new.append(self.box(b[1], b[2]))
            elif k == "fig":
                new.append(self.picture(FIGURES / b[1]))
                new.append(self.fig_caption(f"Figure {b[2]}.", b[3], b[4]))
            elif k == "table":
                new += [self.table_caption(f"Table {b[1]}.", b[2]), self.table(b[3], b[4], b[5]), self.source_line(b[6])]
        anchor = self.tpl[161]                      # 'Appendices' heading
        for e in self.tpl[64:161]:
            self.body.remove(e)
        for e in new:
            anchor.addprevious(e)
        for p, txt in zip(self.tpl[7:12], EXEC):    # executive summary paragraphs
            rs = p.findall(W + "r")
            for r in rs[1:]:
                p.remove(r)
            rs[0].find(W + "t").text = txt

    def contents_lists(self):
        def line(tpl_i, text, page):
            p = copy.deepcopy(self.tpl[tpl_i])
            rs = p.findall(W + "r")
            rs[0].find(W + "t").text = text
            rs[1].find(W + "t").text = "\t" + str(page)
            return p
        toc = [(14 if b[0] == "h1" else 15, b[1]) for b in BODY if b[0] in ("h1", "h2")]
        toc += [(14 if a in ("Appendices", "References") else 15, a) for a in APPENDICES]
        for e in self.tpl[13:44]:
            self.body.remove(e)
        for tpl_i, text in toc:
            self.tpl[44].addprevious(line(tpl_i, text, self.pages.get(text, "")))
        for e in self.tpl[46:58] + self.tpl[59:64]:
            self.body.remove(e)
        for b in BODY:
            if b[0] == "fig":
                self.tpl[58].addprevious(line(46, f"Figure {b[2]}. {b[3]}", self.pages.get(f"Figure {b[2]}.", "")))
        after = self.tpl[58]
        for b in BODY:
            if b[0] == "table":
                e = line(46, f"Table {b[1]}. {b[2]}", self.pages.get(f"Table {b[1]}.", ""))
                after.addnext(e)
                after = e

    # -------------------------------------------------------------- 2. appendices
    def appendices(self):
        cap, src, body_tpl, h2 = self.tpl[125], self.tpl[127], self.tpl[83], self.tpl[162]
        # Table E3: the segment table leaves the main body for Appendix E
        fe1 = self.find("Figure E1.")
        fe1.addnext(self.para_from(src, [("Source: Sections 2.2 and 2.7; World Bank (2025). Sizes and barriers are hypotheses "
                                          "for internal data to test. KPMG Vietnam (2025) expects 68.2 per cent of households "
                                          "to earn more than USD 5,000 a year by 2028.", False, 17)]))
        fe1.addnext(copy.deepcopy(self.tpl[126]))
        fe1.addnext(self.para_from(cap, [("Table E3. ", True, 19), ("Customer segments and objectives", False, 19)]))
        # the scoring-sensitivity table and its anchors are no longer used
        tf2 = next(e for e in self.body.iterchildren() if e.tag == W + "tbl" and "First under 20,000 random weightings" in text_of(e))
        for e in (tf2.getprevious(), tf2.getnext(), tf2, self.find("F1. Scale anchors")):
            self.body.remove(e)
        p = self.find("F2. Scenario method")
        self.replace_in(p, "F2. Scenario method", "F1. Scenario method")
        self.replace_in(p, "Table 4", "Table FX5")
        self.replace_in(p, "(test 6)", "(test 7)")
        p.append(run(" The upper bound on tier interest assumes that every saver keeps one full month's contribution in the "
                     "idle-cash account at 2.5 points above the base rate, the gap between 1.8 and 4.3 per cent. The comparison "
                     "with sign-up rewards multiplies VND 100,000 by the 518,514 accounts opened in 2025. The investor-cash "
                     "comparison in Section 3.2 divides industry investor cash of VND 86.6 trillion by 13.85 million accounts, "
                     "which gives VND 6.25 million per account, against VND 1.15 million at DNSE."))
        self.replace_in(self.find("Figure F1."), "Source: Table 4.", "Source: Table FX5.")
        for e in self.body.iterchildren():
            if e.tag == W + "tbl" and "Segment sizes and barriers (Table 1)" in text_of(e):
                for t in e.iter(W + "t"):
                    if t.text == "Segment sizes and barriers (Table 1)":
                        t.text = "Segment sizes and barriers (Table E3)"
                    elif t.text == "Tests 2 and 3":
                        t.text = "Tests 2 to 4"
                    elif t.text and t.text.startswith("Test 4 (SSC"):
                        t.text = "Test 5 (compliance review and SSC notice)"
        for e in self.appendix_elements():              # template Tables F3 to F5 become F2 to F4
            for t in e.iter(W + "t"):
                if t.text and "Table F" in t.text:
                    t.text = re.sub(r"Table F3\b", "Table F2", t.text)
                    t.text = re.sub(r"Table F4\b", "Table F3", t.text)
                    t.text = re.sub(r"Table F5\b", "Table F4", t.text)
        # Table F5 (scenarios, recomputed from data/manual/scenario_assumptions.csv)
        sc = AN.scenarios()
        hu = lambda x, d: f"{AN_round(x, d)}"
        rows = [["Savers enrolled"] + [f"{s['savers']:,.0f}" for s in sc],
                ["Monthly contribution, VND million"] + [f"{s['monthly_contribution_vnd_m']:.1f}" for s in sc],
                ["Share of contributions kept"] + [f"{s['share_kept'] * 100:.0f}%" for s in sc],
                ["Assets, VND trillion (multiple of investor cash)"] + [f"{hu(s['assets_vnd_tn'], 2)} ({s['multiple_of_investor_cash']:.1f})" for s in sc],
                ["Annual fees, VND billion"] + [hu(s["annual_fees_vnd_bn"], 0) for s in sc],
                ["Tier interest, upper bound, VND billion a year"] + [hu(s["tier_interest_upper_bound_vnd_bn"], 1) for s in sc],
                ["Accounts with VND 10 million or more (share of mid-2026 base)"] +
                [f"{s['accounts_10m_plus']:,.0f} ({s['share_of_mid_2026_base_pct']:.1f}%)" for s in sc]]
        f1p = self.find("F1. Scenario method")
        f1p.addprevious(self.para_from(cap, [("Table F5. ", True, 19), ("Outcomes 18 months after launch (all inputs are assumptions)", False, 19)]))
        f1p.addprevious(self.table([""] + [s["scenario"] for s in sc], rows, [4312, 1700, 1700, 1700], align_numbers=False))
        f1p.addprevious(self.para_from(src, [("Source: authors' projections (Appendix F1; Figure F1).", False, 17)]))
        # validation direction, Table F6 and Figure F2
        val = [["1", "Many dormant accounts belong to salaried intenders", "Internal data; survey of 2,000 customers", "30% or more intend to invest"],
               ["2", "Advised automation turns intention into funding", "Randomised pilot with 20,000 customers", "15% enrol in 60 days; 70% of money kept at day 90"],
               ["3", "Tiers and the skip month raise persistence", "Pilot arms with and without tiers", "Six-contribution rate up 10 points"],
               ["4", "Trust features raise set-up", "A/B test of the trust screen", "Set-up completion up 5 points"],
               ["5", "The advice flow is compliant (gate)", "Compliance review; SSC notice", "Written confirmation"],
               ["6", "Transfers are reliable (gate)", "Load tests; shadow payday runs", "99.5% transfer success"],
               ["7", "Inflows are new money; unit costs fall", "Tag inflow sources; track costs", "80% new money; cost below VND 7.3 million per funded account"]]
        new = [self.para_from(body_tpl, [("F2. Validation direction. ", True, None), (
                   "The headline KPI is the number of accounts holding VND 10 million or more, which DNSE already reports. It was "
                   "27,100 (1.8 per cent of accounts) at the end of 2025, and the target is 85,000 within 18 months, 5 per cent of "
                   "the mid-2026 base. Supporting KPIs are the share of savers reaching three and six consecutive contributions, net "
                   "new money, cost per funded account (VND 7.3 million today) and the trust share of negative reviews (37 per cent "
                   "since 2025). Table F6 ranks the assumptions by the damage they would cause if wrong. DNSE should scale only if "
                   "tests 1 to 4 pass and both gates hold. If enrolment falls short while retention holds, the message changes, and "
                   "if retention fails, the concept stops (Figure F2).", False, None)]),
               self.para_from(cap, [("Table F6. ", True, 19), ("Priority assumptions, tests and thresholds for Round 3", False, 19)]),
               self.table(["#", "Assumption", "Test", "Threshold to proceed"], val, [400, 3200, 3100, 2712], align_numbers=False),
               self.para_from(src, [("Source: authors' design.", False, 17)]),
               self.picture(FIGURES / "figF2_roadmap.png"),
               self.fig_caption("Figure F2.", "Validation roadmap for Round 3", "Source: authors' design.")]
        anchor = self.find("Figure F1.")
        for e in new:
            anchor.addnext(e)
            anchor = e
        # Appendix G
        ref_h = self.references_heading()
        g_h = copy.deepcopy(h2)
        rs = g_h.findall(W + "r")
        for r in rs[1:]:
            g_h.remove(r)
        rs[0].find(W + "t").text = "Appendix G. Regulatory basis and technology proposals"
        G = [["Proposal", "What it would do", "Main constraint", "Decision"],
             ["Ensa suitability check and model portfolios", "Records goal, horizon, income and risk tolerance; proposes one of three model portfolios", "This is investment advice; Circular 121/2020, Article 24 requires this client information", "Adopt inside DNSE's advisory licence; the customer confirms"],
             ["Automatic rebalancing", "Sells and buys to restore target weights", "Article 4 bars investing on a client's behalf without a mandate; each sale incurs 0.1% tax", "Adopt with new money only; changes confirmed in one tap"],
             ["Paid membership tier", "Monthly fee for extra features", "Fees before any return repeat the terms that customers call deceptive", "Replace with free tiers earned by contributions"],
             ["ZaloPay round-up micro-investing", "Rounds up wallet payments and invests the difference", "Small amounts; needs ZaloPay's product work; funding of cashback unclear", "Phase-2 test for young users"],
             ["Ensa Smart Vaults via Lightspeed API", "Algorithmic portfolios executed automatically for an asset-based fee", "Portfolio management is outside a securities company's four licensed businesses (Law on Securities, Article 72); speed adds little for monthly savers", "Set aside; possible later as funds from a licensed manager that DNSE distributes"],
             ["Margin-on-Wealth credit line", "Overdraft secured on Golden Egg, fund units and idle cash", "Article 27 limits lending to margin and defined purposes; fund units are not listed securities; 2026 fine on unreported advance services", "Reject"],
             ["Margin Deal discount as a tier reward", "Cheaper leverage after twelve contributions", "Conflicts with a saver's goal and with the 2026 inspection findings", "Exclude from rewards"],
             ["Transparency dashboard and VNeID checks", "Shows the segregated client account; faster identity checks", "Needs a bank data feed; VNeID integration to be confirmed", "Adopt the dashboard; test VNeID"]]
        for e in (g_h,
                  self.para_from(body_tpl, [("Table G1 records how each proposal was tested against the licensing and lending rules "
                                             "that apply to a securities company. The decisions define the scope of the first release in "
                                             "Section 4.", False, None)]),
                  self.para_from(cap, [("Table G1. ", True, 19), ("Assessment of the proposed features and technology add-ons", False, 19)]),
                  self.table(G[0], G[1:], [2000, 2500, 3000, 1912], align_numbers=False),
                  self.para_from(src, [("Source: Ministry of Finance (2020); National Assembly (2019); Market Times (2026); "
                                        "DNSE Securities (2026d); authors' assessment.", False, 17)])):
            ref_h.addprevious(e)
        # new references, in alphabetical order
        ref_tpl = self.find("Báo Đấu thầu.")

        def ref(text):
            p = copy.deepcopy(ref_tpl)
            rs = p.findall(W + "r")
            for r in rs[1:]:
                p.remove(r)
            rs[0].find(W + "t").text = text
            return p
        ref_tpl.addprevious(ref("Ashraf, N., Karlan, D., & Yin, W. (2006). Tying Odysseus to the mast: Evidence from a commitment savings product in the Philippines. The Quarterly Journal of Economics, 121(2), 635-672. https://doi.org/10.1162/qjec.2006.121.2.635"))
        d26c = self.find("DNSE Securities. (2026c)")
        d26e = ref("DNSE Securities. (2026e). Lightspeed API [Product page]. Retrieved September 27, 2026, from https://www.dnse.com.vn/lightspeed-api")
        d26d = ref('DNSE Securities. (2026d). Điều khoản và điều kiện sản phẩm "Ý Tưởng Đầu Tư" [Terms and conditions of the Investment Ideas product]. Retrieved September 27, 2026, from https://hdsd.dnse.com.vn/die-u-khoa-n-di-ch-vu-dnse/dieu-khoan-san-pham-dich-vu/dieu-khoan-va-dieu-kien-san-pham-y-tuong-dau-tu')
        d26c.addnext(d26e)
        d26c.addnext(d26d)
        self.find("Karlan, D.").addprevious(ref("Kahneman, D., & Tversky, A. (1979). Prospect theory: An analysis of decision under risk. Econometrica, 47(2), 263-291. https://doi.org/10.2307/1914185"))
        mt = self.find("Market Times.")
        mt.addnext(ref("National Assembly of Vietnam. (2019). Law on Securities No. 54/2019/QH14. https://thuvienphapluat.vn/van-ban/Chung-khoan/Luat-Chung-khoan-nam-2019-399763.aspx"))
        mt.addnext(ref("Ministry of Finance. (2020). Circular No. 121/2020/TT-BTC on the operation of securities companies. https://thuvienphapluat.vn/van-ban/Doanh-nghiep/Thong-tu-121-2020-TT-BTC-huong-dan-hoat-dong-cua-cong-ty-chung-khoan-453690.aspx"))
        # English product names in the text; English renderings of Vietnamese publication names everywhere
        in_refs = False
        for e in list(self.body.iterchildren()):
            if e is ref_h:
                in_refs = True
            for t in e.iter(W + "t"):
                if not t.text:
                    continue
                x = t.text
                if not in_refs:
                    for a, b in [("Tài khoản không ngủ", "Never-Sleeping Account"), ("Tài khoản Không Ngủ", "Never-Sleeping Account"),
                                 ("Trứng Vàng", "Golden Egg"), ("Ý tưởng đầu tư", "Investment Ideas"), ("Ý Tưởng Đầu Tư", "Investment Ideas")]:
                        x = x.replace(a, b)
                for a, b in [("Nhân Dân", "Nhan Dan"), ("Thanh Niên", "Thanh Nien"), ("Báo Đấu thầu", "Bao Dau Thau"),
                             ("Phụ nữ Việt Nam", "Phu Nu Viet Nam"), ("Thương hiệu & Công luận", "Thuong Hieu & Cong Luan"),
                             ("QĐ-XPHC", "QD-XPHC"), ("Table FX5", "Table F5")]:
                    x = x.replace(a, b)
                t.text = x

    # -------------------------------------------------------------- 3. appendix figures drawn by src/figures
    def swap_appendix_figures(self):
        for caption, png in APPENDIX_FIGURES.items():
            img_p = self.find(caption).getprevious()
            blip = next(img_p.iter(A_NS + "blip"))
            part = self.doc.part.related_parts[blip.get(R_NS + "embed")]
            part._blob = (FIGURES / png).read_bytes()
            cx, cy = PRINT_SIZE[png]
            exts = [e for e in img_p.iter(A_NS + "ext") if e.getparent().tag == A_NS + "xfrm"]
            for ext in list(img_p.iter(WP_NS + "extent")) + exts:
                ext.set("cx", str(cx))
                ext.set("cy", str(cy))

    # -------------------------------------------------------------- 4. references no longer cited
    def prune_references(self):
        els = list(self.body.iterchildren())
        ref_h = self.references_heading()
        i_ref = els.index(ref_h)
        text = "\n".join(text_of(e) for e in els[:i_ref])
        checks = {"FiinRatings. (2026)": [r"FiinRatings"], "Kemp, S. (2026)": [r"Kemp"],
                  "Karlan, D., McConnell": [r"Karlan et al"], "Thanh Nien. (2025": [r"Thanh Nien,? \(?2025"],
                  "Vietstock. (2025": [r"Vietstock,? \(?2025"], "Vietstock. (2026d": [r"Vietstock,? \(?2026d", r"2026[a-e], 2026d"],
                  "Vietstock. (2026e": [r"Vietstock,? \(?2026e"]}
        for e in els[i_ref + 1:]:
            for start, pats in checks.items():
                if text_of(e).startswith(start) and not any(re.search(p, text) for p in pats):
                    self.body.remove(e)
                    print("  dropped uncited reference:", start)

    def drop_unused_images(self):
        used = set(self.body.xpath(".//a:blip/@r:embed"))
        for rid, rel in list(self.doc.part.rels.items()):
            if rel.reltype == RT.IMAGE and rid not in used:
                del self.doc.part.rels[rid]

    def renumber_drawings(self):
        """Give every drawing a unique id (copied template pictures would otherwise share ids)."""
        for i, e in enumerate(self.body.iter(WP_NS + "docPr"), start=1):
            e.set("id", str(i))

    def build(self, out):
        self.main_body()
        self.contents_lists()
        self.appendices()
        self.swap_appendix_figures()
        self.prune_references()
        build_title_page(self.doc)
        self.drop_unused_images()
        self.renumber_drawings()
        self.doc.save(str(out))
        return out


def AN_round(x, d):
    """Half-up rounding after rounding to one more decimal, as used in the report tables."""
    from decimal import Decimal, ROUND_HALF_UP
    q = Decimal(str(round(x, d + 1))).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP)
    return f"{q:,}"


def word_count(path):
    """Words in Sections 1 to 5, including tables, captions and headings."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = Document(str(path))
    on, n = False, 0
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            t = Paragraph(el, d).text
            sty = el.xpath("./w:pPr/w:pStyle/@w:val")
            if t.strip() == "1. Introduction" and sty == ["Heading1"]:
                on = True
            if t.strip() == "Appendices" and sty == ["Heading1"]:
                on = False
            if on:
                n += len(t.split())
        elif tag == "tbl" and on:
            for r in Table(el, d).rows:
                seen = set()
                for c in r.cells:
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc))
                        n += len(c.text.split())
    return n


def soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(p).exists():
            return p
    return None


def to_pdf(docx):
    exe = soffice()
    if not exe:
        print("  LibreOffice not found: PDF skipped (open the .docx in Word and save as PDF instead)")
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return docx.with_suffix(".pdf")


def page_map(pdf):
    from pypdf import PdfReader
    texts = [re.sub(r"\s+", " ", pg.extract_text() or "") for pg in PdfReader(str(pdf)).pages]
    lof = next(i for i, t in enumerate(texts) if "List of figures and tables" in t)
    keys = [b[1] for b in BODY if b[0] in ("h1", "h2")] + APPENDICES
    labels = [f"Figure {b[2]}." for b in BODY if b[0] == "fig"] + [f"Table {b[1]}." for b in BODY if b[0] == "table"]
    pages = {}
    for k in keys:
        for i in range(lof + 1, len(texts)):
            if k[:40] in texts[i]:
                pages[k] = i + 1
                break
    for k in labels:
        for i in range(lof + 1, len(texts)):
            if re.search(r"(^| )" + re.escape(k) + " ", texts[i]):
                pages[k] = i + 1
                break
    return pages


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", action="store_true", help="export a PDF with LibreOffice")
    ap.add_argument("--paginate", action="store_true", help="recompute page numbers with LibreOffice")
    a = ap.parse_args()
    pages = json.loads(PAGES.read_text(encoding="utf-8"))
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    out = Builder(pages).build(OUT_DOCX)
    if a.paginate:
        for _ in range(3):
            pdf = to_pdf(out)
            if not pdf:
                break
            new = page_map(pdf)
            if new == pages:
                break
            pages = new
            PAGES.write_text(json.dumps(pages, ensure_ascii=False, indent=0), encoding="utf-8")
            out = Builder(pages).build(OUT_DOCX)
    if a.pdf or a.paginate:
        to_pdf(out)
    print(f"  wrote {out.relative_to(ROOT)}; main body {word_count(out):,} words")


if __name__ == "__main__":
    main()
